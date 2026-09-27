import os
from PyQt6.QtWidgets import QProgressDialog, QDialog, QMessageBox, QFileDialog, QApplication
from PyQt6.QtCore import Qt
from core.database import Person
from network.client import get_server_url


def _batch_import_ids(self):
    """从数据库按姓名或身份证号匹配，将人员信息填入Excel表格"""
    path, _ = QFileDialog.getOpenFileName(
        self, "选择需要填写数据的Excel文件", "",
        "Excel文件 (*.xlsx)"
    )
    if not path:
        return

    try:
        import openpyxl
        wb = openpyxl.load_workbook(path)
        wb_values = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        QMessageBox.critical(self, "错误", f"读取Excel文件失败：{e}")
        return

    from dialogs.report.column_mapping_dialog import ColumnMappingDialog, MATCH_BY_NAME, MATCH_BY_ID_CARD
    mapping_dialog = ColumnMappingDialog(wb, self)
    if mapping_dialog.exec() != QDialog.DialogCode.Accepted:
        return

    match_by = mapping_dialog.match_by
    match_col = mapping_dialog.match_col
    fill_map = mapping_dialog.fill_map
    header_row = mapping_dialog.header_row
    ws = mapping_dialog.ws
    sheet_name = ws.title
    ws_values = wb_values[sheet_name]

    if not fill_map:
        QMessageBox.warning(self, "提示", "请至少选择一列需要填写的数据")
        return

    # 一次性加载全部人员数据到内存，本地匹配（宏查询模式）
    progress = QProgressDialog("正在加载数据库人员数据...", "取消", 0, 0, self)
    progress.setWindowTitle("批量填写")
    progress.setWindowModality(Qt.WindowModality.WindowModal)
    progress.setMinimumDuration(0)
    progress.setValue(0)
    QApplication.processEvents()

    all_persons = []
    if get_server_url():
        from network.client import api_persons_lookup_all
        try:
            all_persons = api_persons_lookup_all()
        except Exception as e:
            QMessageBox.critical(self, "错误", f"加载服务器数据失败：{e}")
            return
    else:
        from core.database import get_session, Person
        session = get_session()
        try:
            for p in session.query(Person).all():
                all_persons.append({
                    'name': p.name or '', 'id_card': p.id_card or '',
                    'phone': p.phone or '', 'social_card1': p.social_card1 or '',
                    'social_card2': p.social_card2 or '', 'social_card3': p.social_card3 or '',
                    'hu_hao': p.hu_hao or '', 'address': p.address or '',
                    'relation': p.relation or ''
                })
        finally:
            session.close()

    progress.cancel()

    if not all_persons:
        QMessageBox.warning(self, "提示", "数据库中没有人员数据")
        return

    # 构建查找索引：姓名→人员列表，身份证号→人员
    name_index = {}
    id_card_index = {}
    for p in all_persons:
        if p['name']:
            name_index.setdefault(p['name'], []).append(p)
        if p['id_card']:
            id_card_index[p['id_card']] = p

    total = ws.max_row - header_row
    if total <= 0:
        QMessageBox.warning(self, "提示", "文件中没有数据行")
        return

    progress = QProgressDialog("正在处理...", "取消", 0, total, self)
    progress.setWindowTitle("批量填写进度")
    progress.setWindowModality(Qt.WindowModality.WindowModal)
    progress.setMinimumDuration(0)

    from openpyxl.styles import Font as XlFont
    red_font = XlFont(color="FF0000")

    found = 0
    not_found = 0
    multi_selected = 0
    skipped = 0
    not_found_values = []

    formula_count = 0

    for row_idx in range(header_row + 1, ws.max_row + 1):
        if progress.wasCanceled():
            break

        # 双重读取：先试缓存值(data_only)，不可用则回退到普通模式
        raw_val = ws_values.cell(row_idx, match_col).value
        if raw_val is None or (isinstance(raw_val, str) and raw_val.startswith('#')):
            # 缓存值为None或错误值(#N/A等)，回退到普通模式
            raw_val = ws.cell(row_idx, match_col).value
        if isinstance(raw_val, str) and raw_val.startswith('='):
            # 公式无法计算，跳过
            formula_count += 1
            progress.setValue(row_idx - header_row)
            continue

        match_val = str(raw_val).strip() if raw_val is not None else ''
        if not match_val:
            skipped += 1
            progress.setValue(row_idx - header_row)
            continue

        progress.setLabelText(f"第 {row_idx - header_row}/{total} 行: {match_val}")
        QApplication.processEvents()

        # 本地匹配，零网络请求
        results = []
        if match_by == MATCH_BY_NAME:
            results = name_index.get(match_val, [])
        else:
            p = id_card_index.get(match_val)
            if p:
                results = [p]

        if len(results) == 0:
            not_found += 1
            not_found_values.append(match_val)
            for col_idx in fill_map:
                cell = ws.cell(row_idx, col_idx, "未查询到")
                cell.font = red_font
        elif len(results) == 1:
            found += 1
            person = results[0]
            for col_idx, field in fill_map.items():
                ws.cell(row_idx, col_idx, person.get(field, ''))
        else:
            from dialogs.report.multi_name_select_dialog import MultiNameSelectDialog
            select_dialog = MultiNameSelectDialog(results, match_val, self)
            if select_dialog.exec() == QDialog.DialogCode.Accepted and select_dialog.selected_person:
                person = select_dialog.selected_person
                found += 1
                multi_selected += 1
                for col_idx, field in fill_map.items():
                    ws.cell(row_idx, col_idx, person.get(field, ''))
            else:
                skipped += 1

        progress.setValue(row_idx - header_row)

    progress.setValue(total)

    # 保存前把匹配列的公式单元格转为静态值，防止Excel重新计算时外部引用失效变#N/A
    for row_idx in range(header_row + 1, ws.max_row + 1):
        cell = ws.cell(row_idx, match_col)
        if isinstance(cell.value, str) and cell.value.startswith('='):
            cached = ws_values.cell(row_idx, match_col).value
            if cached is not None and not (isinstance(cached, str) and cached.startswith('#')):
                cell.value = cached

    default_name = os.path.splitext(os.path.basename(path))[0] + "_已填写.xlsx"
    save_path, _ = QFileDialog.getSaveFileName(
        self, "保存填写后的文件", default_name,
        "Excel文件 (*.xlsx)",
        options=QFileDialog.Option.DontUseNativeDialog
    )
    if not save_path:
        return

    try:
        wb.save(save_path)
    except Exception as e:
        QMessageBox.critical(self, "保存失败", f"文件保存失败：{e}")
        return

    match_label = "姓名" if match_by == MATCH_BY_NAME else "身份证号"
    summary = f"处理完成！\n\n"
    summary += f"匹配方式: 按{match_label}匹配\n"
    summary += f"总行数: {total}\n"
    summary += f"成功填写: {found} 条（其中手动选择: {multi_selected} 条）\n"
    summary += f"未查询到: {not_found} 条\n"
    summary += f"跳过: {skipped} 条"
    if formula_count:
        summary += f"\n公式无法读取: {formula_count} 条（请在Excel中打开文件让公式计算后再导入）"
    if not_found_values:
        summary += f"\n\n未查询到的{match_label}:\n"
        for v in not_found_values:
            summary += f"  - {v}\n"
    summary += f"\n文件已保存至: {save_path}"

    QMessageBox.information(self, "处理结果", summary)
