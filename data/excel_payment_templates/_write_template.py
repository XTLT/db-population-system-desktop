import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment


def _write_template(path, sheet_title, columns, col_widths=None,
                    title=None, desc=None, label_row=None, header_row=1):
    """通用模板写入：表头样式 + 列宽。

    - 默认单表头行（header_row=1）。
    - 4 行式花名册：传入 title（第1行标题）、desc（第2行填表说明）、
      label_row（dict: 1-based 列号 -> 标签文字，写在表头行上一行），
      header_row=4 时表头从第4行开始写。
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title

    n = len(columns)

    if title:
        cell = ws.cell(row=1, column=1, value=title)
        cell.alignment = Alignment(horizontal="center")
    if desc:
        cell = ws.cell(row=2, column=1, value=desc)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    if label_row:
        for col_idx, text in label_row.items():
            cell = ws.cell(row=header_row - 1, column=col_idx, value=text)
            cell.alignment = Alignment(horizontal="center")

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    if col_widths:
        for col_idx, width in enumerate(col_widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width
    else:
        for col_idx in range(1, n + 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 15

    wb.save(path)
