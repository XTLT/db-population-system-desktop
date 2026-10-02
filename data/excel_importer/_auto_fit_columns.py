import openpyxl
from openpyxl.utils import get_column_letter


def _auto_fit_columns(file_path: str, min_width: float = 10, max_width: float = 40):
    """按单元格内容自动调整 Excel 列宽（中文按 2 个字符宽度估算）。

    解决 pandas/DataFrame 导出后列宽默认 8.4 字符导致身份证号、
    地址、备注、社保卡号等长文本显示不全的问题。
    """
    try:
        wb = openpyxl.load_workbook(file_path)
        for ws in wb.worksheets:
            for col_cells in ws.columns:
                col_letter = get_column_letter(col_cells[0].column)
                max_len = min_width
                for cell in col_cells:
                    if cell.value is None:
                        continue
                    s = str(cell.value)
                    length = sum(2 if ord(ch) > 127 else 1 for ch in s)
                    if length > max_len:
                        max_len = length
                ws.column_dimensions[col_letter].width = min(max_len + 2, max_width)
        wb.save(file_path)
    except Exception:
        # 列宽调整失败不影响导出结果
        pass
